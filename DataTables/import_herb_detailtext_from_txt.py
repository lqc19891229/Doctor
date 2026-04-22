# -*- coding: utf-8 -*-
"""
用途：
1. 读取当前脚本同目录下的 Data.xlsx
2. 读取当前脚本同目录下【原文】文件夹中的 .txt 文件
3. 以 txt 文件名 = HerbName 为匹配基准
4. 将 txt 文件内容写入 Herb sheet 的 DetailText 列
5. 覆盖保存 Data.xlsx，并额外生成一个备份文件

运行方式：
    在 Doctor/DataTables 目录下执行：
        python import_herb_detailtext_from_txt.py

依赖：
    pip install openpyxl
"""

from __future__ import annotations

from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook


# 当前脚本所在目录（预期为 Doctor/DataTables）
CURRENT_DIR = Path(__file__).resolve().parent
# Excel 文件路径
EXCEL_PATH = CURRENT_DIR / "Data.xlsx"
# 原文 txt 文件所在目录
SOURCE_DIR = CURRENT_DIR / "原文"
# 备份文件路径
BACKUP_PATH = CURRENT_DIR / "Data.backup_before_detailtext_import.xlsx"
# 目标 sheet 名称
SHEET_NAME = "Herb"


def log(message: str) -> None:
    """统一日志输出。"""
    print(f"[import_herb_detailtext] {message}")


def normalize_text(text: str) -> str:
    """
    清洗 txt 内容：
    1. 去掉 UTF-8 BOM。
    2. 统一换行符为 \n。
    3. 去掉首尾多余空白，但保留正文内部换行。
    """
    text = text.replace("\ufeff", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def load_txt_map(source_dir: Path) -> dict[str, str]:
    """
    读取【原文】目录下所有 txt 文件，构建：
        文件名（不带扩展名） -> 文件内容
    例如：
        生姜.txt -> "这里是生姜原文内容"
    """
    if not source_dir.exists():
        raise FileNotFoundError(f"找不到原文目录：{source_dir}")

    txt_files = sorted(source_dir.glob("*.txt"))
    if not txt_files:
        raise FileNotFoundError(f"原文目录下没有找到任何 .txt 文件：{source_dir}")

    txt_map: dict[str, str] = {}

    for txt_path in txt_files:
        herb_name = txt_path.stem.strip()
        content = normalize_text(txt_path.read_text(encoding="utf-8"))

        if not herb_name:
            log(f"跳过空文件名：{txt_path.name}")
            continue

        txt_map[herb_name] = content

    return txt_map


def find_column_indexes(ws) -> tuple[int, int]:
    """
    查找 Herb sheet 中：
    1. HerbName 列
    2. DetailText 列

    返回值为 openpyxl 使用的 1-based 列索引。
    """
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    herb_name_col = None
    detail_text_col = None

    for index, value in enumerate(header_row, start=1):
        if value == "HerbName":
            herb_name_col = index
        elif value == "DetailText":
            detail_text_col = index

    if herb_name_col is None:
        raise ValueError(f"{SHEET_NAME} sheet 缺少 HerbName 列")
    if detail_text_col is None:
        raise ValueError(f"{SHEET_NAME} sheet 缺少 DetailText 列")

    return herb_name_col, detail_text_col


def import_detail_text() -> None:
    """执行导入逻辑。"""
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"找不到 Excel 文件：{EXCEL_PATH}")

    txt_map = load_txt_map(SOURCE_DIR)
    log(f"已读取 txt 文件数量：{len(txt_map)}")

    # 先备份 Excel，避免覆盖后无法回退
    copy2(EXCEL_PATH, BACKUP_PATH)
    log(f"已创建备份：{BACKUP_PATH.name}")

    wb = load_workbook(EXCEL_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Excel 中不存在 sheet：{SHEET_NAME}")

    ws = wb[SHEET_NAME]
    herb_name_col, detail_text_col = find_column_indexes(ws)

    matched_count = 0
    missing_txt_names: list[str] = []
    unused_txt_names = set(txt_map.keys())

    # 从第 2 行开始遍历数据行，第 1 行为表头
    for row_index in range(2, ws.max_row + 1):
        herb_name = ws.cell(row=row_index, column=herb_name_col).value
        herb_name = str(herb_name).strip() if herb_name is not None else ""

        if not herb_name:
            continue

        if herb_name in txt_map:
            ws.cell(row=row_index, column=detail_text_col).value = txt_map[herb_name]
            matched_count += 1
            unused_txt_names.discard(herb_name)
        else:
            missing_txt_names.append(herb_name)

    wb.save(EXCEL_PATH)

    log(f"写入完成，成功匹配：{matched_count} 条")

    if missing_txt_names:
        log("以下 HerbName 没有找到同名 txt，未写入：")
        for name in missing_txt_names:
            print(f" - {name}")

    if unused_txt_names:
        log("以下 txt 没有匹配到 Herb sheet 中的 HerbName：")
        for name in sorted(unused_txt_names):
            print(f" - {name}")

    if not missing_txt_names and not unused_txt_names:
        log("全部匹配成功")


if __name__ == "__main__":
    import_detail_text()
