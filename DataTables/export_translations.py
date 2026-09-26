#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import re
import sys
from copy import copy
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误：缺少 openpyxl 依赖。")
    print("请先安装：pip install openpyxl")
    raise SystemExit(1)


SHEET_ORDER = ["UI", "Herbs", "Formulas", "Diseases", "NPC", "Story"]
EXPECTED_HEADER = ["keys", "zh_CN", "en", "ja"]
SEARCH_HEADER = ["keys", "zh_CN", "ja_kana", "ja_romaji"]
SEARCH_EXPORT_HEADER = ["keys", "ja_kana", "ja_romaji"]

SYMPTOM_KEY_PREFIX = "UI_DISEASE_SYMPTOM_"
PLACEHOLDER_RE = re.compile(r"%(?:[-+0 #]*\d*(?:\.\d+)?)?[sdif]")


def as_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def placeholders(value: str) -> list[str]:
    return PLACEHOLDER_RE.findall(value or "")


def find_data_xlsx(explicit_path: Path | None, managed_xlsx: Path) -> Path | None:
    """
    Locate Data.xlsx.

    Priority:
    1. Third command-line argument
    2. Same folder as translations_managed.xlsx
    3. ../DataTables/Data.xlsx relative to translations_managed.xlsx
    4. ./DataTables/Data.xlsx relative to translations_managed.xlsx
    5. Current working directory variants
    """
    if explicit_path is not None:
        return explicit_path if explicit_path.exists() else None

    candidates = [
        managed_xlsx.parent / "Data.xlsx",
        managed_xlsx.parent.parent / "DataTables" / "Data.xlsx",
        managed_xlsx.parent / "DataTables" / "Data.xlsx",
        Path("Data.xlsx"),
        Path("../DataTables/Data.xlsx"),
        Path("DataTables/Data.xlsx"),
    ]

    seen = set()
    for path in candidates:
        try:
            resolved = path.resolve()
        except OSError:
            resolved = path

        key = str(resolved)
        if key in seen:
            continue
        seen.add(key)

        if path.exists():
            return path

    return None


def validate_managed_sheets(wb) -> None:
    missing = [name for name in SHEET_ORDER if name != "Story" and name not in wb.sheetnames]
    if missing:
        raise ValueError("翻译工作簿缺少以下工作表： " + ", ".join(missing))

    for sheet_name in SHEET_ORDER:
        if sheet_name == "Story" and sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [as_text(ws.cell(1, col).value) for col in range(1, 5)]
        if header != EXPECTED_HEADER:
            raise ValueError(
                f"{sheet_name} 的表头必须是 {EXPECTED_HEADER}，当前实际为 {header}"
            )

    if "JapaneseSearch" not in wb.sheetnames:
        raise ValueError("翻译工作簿缺少 JapaneseSearch 工作表")
    search_header = [as_text(wb["JapaneseSearch"].cell(1, col).value) for col in range(1, 5)]
    if search_header != SEARCH_HEADER:
        raise ValueError(f"JapaneseSearch 的表头必须是 {SEARCH_HEADER}")


def read_disease_symptoms(data_xlsx: Path) -> list[tuple[str, list[str]]]:
    """
    Read Data.xlsx -> Disease -> DiseaseID + Symptoms.
    Symptoms uses | as the separator.
    """
    wb = load_workbook(data_xlsx, data_only=True, read_only=True)

    if "Disease" not in wb.sheetnames:
        raise ValueError(f"{data_xlsx} 缺少 Disease 工作表")

    ws = wb["Disease"]
    headers = {
        as_text(cell.value): index
        for index, cell in enumerate(ws[1], start=1)
        if as_text(cell.value)
    }

    for required in ("DiseaseID", "Symptoms"):
        if required not in headers:
            raise ValueError(f"Data.xlsx 的 Disease 工作表缺少列：{required}")

    id_col = headers["DiseaseID"]
    symptoms_col = headers["Symptoms"]

    result: list[tuple[str, list[str]]] = []
    seen_ids: set[str] = set()

    for row_no in range(2, ws.max_row + 1):
        disease_id = as_text(ws.cell(row_no, id_col).value)
        raw_symptoms = as_text(ws.cell(row_no, symptoms_col).value)

        if not disease_id:
            continue

        if disease_id in seen_ids:
            raise ValueError(
                f"Data.xlsx 的 Disease 工作表第 {row_no} 行存在重复的 DiseaseID： "
                f"{disease_id}"
            )
        seen_ids.add(disease_id)

        symptoms = [
            item.strip()
            for item in raw_symptoms.split("|")
            if item and item.strip()
        ]

        result.append((disease_id, symptoms))

    return result


def parse_symptom_key(key: str) -> tuple[str, int] | None:
    if not key.startswith(SYMPTOM_KEY_PREFIX):
        return None

    rest = key[len(SYMPTOM_KEY_PREFIX):]
    match = re.fullmatch(r"(.+)_([0-9]{2,})", rest)
    if match is None:
        return None

    return match.group(1).lower(), int(match.group(2))


def copy_row_style(source_ws, source_row: int, target_ws, target_row: int) -> None:
    for col in range(1, 5):
        source = source_ws.cell(source_row, col)
        target = target_ws.cell(target_row, col)

        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)

    target_ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height


def update_table_ranges(ws) -> None:
    """
    The managed workbook uses Excel tables.
    After deleting/appending symptom rows, extend any table on this sheet to the new end row.
    """
    if ws.max_row < 1:
        return

    for table in ws.tables.values():
        table.ref = f"A1:D{ws.max_row}"


def sync_disease_symptoms(managed_wb, disease_symptoms: list[tuple[str, list[str]]]) -> list[str]:
    """
    Synchronize UI_DISEASE_SYMPTOM_* rows in the Diseases sheet.

    Translation preservation priority:
    1. Same disease + same Chinese symptom
    2. Same Chinese symptom anywhere in old symptom rows
    3. Otherwise English remains blank and export stops so it can be translated

    This means reordering Symptoms in Data.xlsx will not attach the wrong English
    translation to the new _01 / _02 / _03 numbering.
    """
    ws = managed_wb["Diseases"]

    existing_by_disease_and_zh: dict[tuple[str, str], tuple[str, str]] = {}
    existing_by_zh: dict[str, tuple[str, str]] = {}
    symptom_rows: list[int] = []

    style_source_row: int | None = None

    for row_no in range(2, ws.max_row + 1):
        key = as_text(ws.cell(row_no, 1).value)
        parsed = parse_symptom_key(key)
        if parsed is None:
            continue

        disease_id, _index = parsed
        zh = as_text(ws.cell(row_no, 2).value)
        en = as_text(ws.cell(row_no, 3).value)
        ja = as_text(ws.cell(row_no, 4).value)

        symptom_rows.append(row_no)

        if style_source_row is None:
            style_source_row = row_no

        if zh:
            existing_by_disease_and_zh[(disease_id, zh)] = (en, ja)
            existing_by_zh.setdefault(zh, (en, ja))

    # If the workbook has no symptom rows yet, use the last ordinary Diseases row
    # as the formatting template.
    if style_source_row is None:
        style_source_row = max(ws.max_row, 2)

    # Remove old symptom rows bottom-up.
    for row_no in reversed(symptom_rows):
        ws.delete_rows(row_no, 1)

    missing_english: list[str] = []

    for disease_id, symptoms in disease_symptoms:
        for index, zh in enumerate(symptoms, start=1):
            key = f"{SYMPTOM_KEY_PREFIX}{disease_id.upper()}_{index:02d}"

            en, ja = existing_by_disease_and_zh.get(
                (disease_id.lower(), zh), existing_by_zh.get(zh, ("", ""))
            )

            target_row = ws.max_row + 1
            copy_row_style(ws, style_source_row, ws, target_row)

            ws.cell(target_row, 1).value = key
            ws.cell(target_row, 2).value = zh
            ws.cell(target_row, 3).value = en
            ws.cell(target_row, 4).value = ja

            if not en:
                missing_english.append(f"{key} | {zh}")

    update_table_ranges(ws)
    return missing_english


def sync_story(wb, data_xlsx: Path) -> list[str]:
    """Rebuild story rows from stable StoryID/LineIndex while retaining reviewed English."""
    source = load_workbook(data_xlsx, read_only=True, data_only=True)
    if "StoryLine" not in source.sheetnames:
        raise ValueError(f"{data_xlsx} 缺少 StoryLine 工作表")
    ws = wb["Story"] if "Story" in wb else wb.create_sheet("Story")
    if [as_text(ws.cell(1, col).value) for col in range(1, 5)] != EXPECTED_HEADER:
        ws.insert_rows(1)
        for col, name in enumerate(EXPECTED_HEADER, start=1):
            ws.cell(1, col).value = name
    old = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            old[str(row[0])] = (as_text(row[1]), as_text(row[2]), as_text(row[3]))
    source_ws = source["StoryLine"]
    headers = {as_text(c.value): i for i, c in enumerate(source_ws[1])}
    required = ("StoryID", "LineIndex", "Speaker", "Text")
    if any(name not in headers for name in required):
        raise ValueError("Data.xlsx 的 StoryLine 工作表缺少必要列")
    rows = []
    seen = set()
    for row in source_ws.iter_rows(min_row=2, values_only=True):
        story_id = as_text(row[headers["StoryID"]])
        if not story_id:
            continue
        index = int(row[headers["LineIndex"]])
        if index < 1 or (story_id, index) in seen:
            raise ValueError(f"StoryLine 行号无效或重复：{story_id} / {index}")
        seen.add((story_id, index))
        for field, suffix in (("Speaker", "SPEAKER"), ("Text", "TEXT")):
            zh = as_text(row[headers[field]])
            if not zh:
                continue
            key = f"STORY_{story_id}_{index:03d}_{suffix}"
            previous_zh, previous_en, previous_ja = old.get(key, ("", "", ""))
            en = previous_en if previous_zh == zh else ""
            ja = previous_ja if previous_zh == zh else ""
            rows.append((key, zh, en, ja))
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for key, zh, en, ja in rows:
        ws.append((key, zh, en, ja))
    update_table_ranges(ws)
    return [f"{key} | {zh}" for key, zh, en, ja in rows if not en]


def save_managed_workbook_safely(wb, path: Path) -> None:
    """
    Save through a temporary file, then replace the original.
    This reduces the risk of leaving a damaged workbook if saving is interrupted.
    """
    temp_path = path.with_name(path.stem + ".__tmp__" + path.suffix)

    try:
        wb.save(temp_path)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def export_csv(wb, out_csv: Path) -> tuple[int, list[str]]:
    all_rows: list[list[str]] = []
    seen: dict[str, str] = {}
    warnings: list[str] = []

    for sheet_name in SHEET_ORDER:
        ws = wb[sheet_name]
        count = 0

        for row_no in range(2, ws.max_row + 1):
            key = as_text(ws.cell(row_no, 1).value)
            zh = as_text(ws.cell(row_no, 2).value)
            en = as_text(ws.cell(row_no, 3).value)

            ja = as_text(ws.cell(row_no, 4).value)

            if not key and not zh and not en and not ja:
                continue

            if not key or not zh or (not en and sheet_name != "Story"):
                raise ValueError(
                    f"数据不完整：{sheet_name} 第 {row_no} 行 "
                    f"(key={key!r}, zh_CN={zh!r}, en={en!r})"
                )

            if key in seen:
                raise ValueError(
                    f"发现重复 key {key!r}：{seen[key]} 与 "
                    f"{sheet_name} row {row_no}"
                )

            seen[key] = f"{sheet_name} row {row_no}"

            zh_ph = placeholders(zh)
            en_ph = placeholders(en)
            if sorted(zh_ph) != sorted(en_ph):
                warnings.append(
                    f"{sheet_name} row {row_no} {key}: "
                    f"zh={zh_ph}, en={en_ph}"
                )

            if ja and sorted(placeholders(ja)) != sorted(zh_ph):
                warnings.append(f"{sheet_name} row {row_no} {key}: ja={placeholders(ja)}, zh={zh_ph}")

            all_rows.append([key, zh, en or zh, ja or zh])
            count += 1

        print(f"{sheet_name}：{count} 条")

    out_csv.parent.mkdir(parents=True, exist_ok=True)

    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(EXPECTED_HEADER)
        writer.writerows(all_rows)

    return len(all_rows), warnings


def export_japanese_search(wb, out_csv: Path) -> int:
    """Export explicit Japanese readings; never derive them from Chinese pinyin IDs."""
    title_keys = {
        as_text(row[0])
        for name in SHEET_ORDER
        for row in wb[name].iter_rows(min_row=2, values_only=True)
        if as_text(row[0]).startswith("UI_BOOK_ENTRY_TITLE_")
    }
    rows = []
    seen = set()
    for row_no, row in enumerate(wb["JapaneseSearch"].iter_rows(min_row=2, values_only=True), start=2):
        key, _source_zh, kana, romaji = (as_text(value) for value in row[:4])
        if not key and not kana and not romaji:
            continue
        if key not in title_keys or key in seen:
            raise ValueError(f"JapaneseSearch 第 {row_no} 行的 key 无效或重复：{key}")
        seen.add(key)
        rows.append((key, kana, romaji))
    if seen != title_keys:
        raise ValueError(f"JapaneseSearch 缺少 {len(title_keys - seen)} 个名称 key")
    # .txt keeps this runtime lookup file as a regular text resource in Godot exports.
    path = out_csv.with_name("search_ja.txt")
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(SEARCH_EXPORT_HEADER)
        writer.writerows(rows)
    return len(rows)


def sync_japanese_search(wb) -> None:
    """Keep search keys in step with entry titles without losing edited readings."""
    ws = wb["JapaneseSearch"]
    existing = {
        as_text(row[0]): (as_text(row[2]), as_text(row[3]))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if as_text(row[0])
    }
    titles = []
    seen = set()
    for name in SHEET_ORDER:
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            key = as_text(row[0])
            if key.startswith("UI_BOOK_ENTRY_TITLE_") and key not in seen:
                seen.add(key)
                titles.append((key, as_text(row[1])))
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for key, zh in titles:
        kana, romaji = existing.get(key, ("", ""))
        ws.append((key, zh, kana, romaji))


def main() -> int:
    managed_xlsx = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path("translations_managed.xlsx")
    )
    out_csv = (
        Path(sys.argv[2])
        if len(sys.argv) > 2
        else Path("translations.csv")
    )
    explicit_data_xlsx = Path(sys.argv[3]) if len(sys.argv) > 3 else None

    if not managed_xlsx.exists():
        print("错误：找不到翻译管理表：", managed_xlsx)
        return 1

    data_xlsx = find_data_xlsx(explicit_data_xlsx, managed_xlsx)
    if data_xlsx is None:
        print("错误：找不到 Data.xlsx。")
        print("请把 Data.xlsx 路径作为第三个参数传入，例如：")
        print(
            "  python export_translations.py "
            "translations_managed.xlsx translations.csv ../DataTables/Data.xlsx"
        )
        return 1

    print("翻译管理表：", managed_xlsx)
    print("疾病数据来源：", data_xlsx)
    print("输出 CSV：", out_csv)
    print()

    try:
        wb = load_workbook(managed_xlsx)
        validate_managed_sheets(wb)

        disease_symptoms = read_disease_symptoms(data_xlsx)
        symptom_count = sum(len(items) for _disease_id, items in disease_symptoms)

        print(
            f"正在同步 Disease!Symptoms： "
            f"{len(disease_symptoms)} 个疾病，{symptom_count} 条症状"
        )

        missing_english = sync_disease_symptoms(wb, disease_symptoms)
        missing_story = sync_story(wb, data_xlsx)
        sync_japanese_search(wb)

        # Always save the synchronized managed workbook first.
        save_managed_workbook_safely(wb, managed_xlsx)
        print("症状已同步到 Diseases 工作表。")

        # New Chinese symptoms need human translation before producing a game CSV.
        if missing_english:
            print()
            print("停止导出：发现新增症状尚未填写英文翻译。")
            print(
                "以下症状已写入 translations_managed.xlsx，但 en 列仍为空："
            )
            for item in missing_english:
                print(" -", item)
            print()
            print("请补全这些英文翻译，保存工作簿后再重新运行本脚本。")
            return 2

        total, warnings = export_csv(wb, out_csv)
        search_count = export_japanese_search(wb, out_csv)
        if missing_story:
            print(f"剧情待翻译：{len(missing_story)} 条；英文暂用中文原文，补全 Story 工作表后重导出。")

        print()
        print(f"已导出 {total} 条翻译 -> {out_csv}")
        print(f"日语搜索读音：{search_count} 条 -> {out_csv.with_name('search_ja.txt')}")

        if warnings:
            print()
            print("警告：发现中英文占位符不一致：")
            for item in warnings:
                print(" -", item)
        else:
            print("占位符检查：通过")

        print()
        print("下一步：在 Godot 中重新导入 Localization/translations.csv。")
        return 0

    except PermissionError as exc:
        print("错误：", exc)
        print(
            "请先关闭 Excel/WPS 中打开的 translations_managed.xlsx，然后重新运行导出脚本。"
        )
        return 1
    except Exception as exc:
        print("错误：", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
