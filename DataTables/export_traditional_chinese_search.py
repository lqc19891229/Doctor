#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误：缺少 openpyxl 依赖。")
    print("请先安装：pip install openpyxl")
    raise SystemExit(1)


SHEET_ORDER = ["UI", "Herbs", "Formulas", "Diseases", "NPC", "Story"]
SEARCH_SHEET = "TraditionalChineseSearch"
SEARCH_HEADER = ["keys", "zh_TW", "zhuyin", "pinyin"]
SEARCH_EXPORT_HEADER = ["keys", "zhuyin", "pinyin"]

SEARCH_IGNORED_CHARS = set(
    " \t\r\n"
    ",，.。・･、"
    "()（）[]【】"
    "/／-—_"
    "·:：;；"
)


def as_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def segmented_parts(value: str, label: str) -> list[str]:
    clean = as_text(value)
    if not clean:
        return []
    if clean.startswith("|") or clean.endswith("|") or "||" in clean:
        raise ValueError(f"{label} 分段格式错误：{clean!r}")
    parts = [part.strip() for part in clean.split("|")]
    if any(not part for part in parts):
        raise ValueError(f"{label} 存在空分段：{clean!r}")
    return parts


def searchable_character_count(value: str) -> int:
    return sum(
        1
        for char in as_text(value)
        if char not in SEARCH_IGNORED_CHARS
    )


def collect_traditional_titles(wb) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    seen: set[str] = set()

    for sheet_name in SHEET_ORDER:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            key = as_text(row[0])
            if not key.startswith("UI_BOOK_ENTRY_TITLE_") or key in seen:
                continue
            zh_tw = as_text(row[2]) if len(row) > 2 else ""
            if not zh_tw:
                raise ValueError(f"{sheet_name} 中 {key} 缺少 zh_TW 标题")
            seen.add(key)
            rows.append((key, zh_tw))

    return rows


def ensure_search_sheet(wb):
    if SEARCH_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(SEARCH_SHEET)
        ws.append(SEARCH_HEADER)
        return ws

    ws = wb[SEARCH_SHEET]
    header = [as_text(ws.cell(1, col).value) for col in range(1, 5)]
    if header != SEARCH_HEADER:
        raise ValueError(
            f"{SEARCH_SHEET} 的表头必须是 {SEARCH_HEADER}，当前实际为 {header}"
        )
    return ws


def sync_search_sheet(wb) -> list[str]:
    ws = ensure_search_sheet(wb)
    existing = {
        as_text(row[0]): (as_text(row[2]), as_text(row[3]))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if as_text(row[0])
    }

    titles = collect_traditional_titles(wb)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    missing: list[str] = []
    for key, zh_tw in titles:
        zhuyin, pinyin = existing.get(key, ("", ""))
        ws.append((key, zh_tw, zhuyin, pinyin))
        if not zhuyin or not pinyin:
            missing.append(f"{key} | {zh_tw}")

    return missing


def validate_and_export(wb, out_path: Path) -> int:
    ws = wb[SEARCH_SHEET]
    title_rows = collect_traditional_titles(wb)
    title_map = dict(title_rows)
    title_keys = set(title_map)

    rows: list[tuple[str, str, str]] = []
    seen: set[str] = set()

    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        key, zh_tw, zhuyin, pinyin = (as_text(v) for v in row[:4])
        if not key and not zh_tw and not zhuyin and not pinyin:
            continue

        if key not in title_keys:
            raise ValueError(f"{SEARCH_SHEET} 第 {row_no} 行存在未知 key：{key}")
        if key in seen:
            raise ValueError(f"{SEARCH_SHEET} 第 {row_no} 行存在重复 key：{key}")
        if not zhuyin or not pinyin:
            raise ValueError(
                f"{SEARCH_SHEET} 第 {row_no} 行缺少注音或拼音：{key}"
            )

        expected_title = title_map[key]
        if zh_tw != expected_title:
            raise ValueError(
                f"{SEARCH_SHEET} 第 {row_no} 行繁体标题不同步："
                f"{key} | 表内={zh_tw!r} | 翻译={expected_title!r}"
            )

        unit_count = searchable_character_count(zh_tw)
        zhuyin_parts = segmented_parts(
            zhuyin, f"{SEARCH_SHEET} 第 {row_no} 行注音"
        )
        pinyin_parts = segmented_parts(
            pinyin, f"{SEARCH_SHEET} 第 {row_no} 行拼音"
        )

        if len(zhuyin_parts) != unit_count:
            raise ValueError(
                f"{SEARCH_SHEET} 第 {row_no} 行不是逐字注音："
                f"{key} | zh_TW={zh_tw!r} ({unit_count} 字) | "
                f"zhuyin={zhuyin!r} ({len(zhuyin_parts)} 段)"
            )
        if len(pinyin_parts) != unit_count:
            raise ValueError(
                f"{SEARCH_SHEET} 第 {row_no} 行不是逐字拼音："
                f"{key} | zh_TW={zh_tw!r} ({unit_count} 字) | "
                f"pinyin={pinyin!r} ({len(pinyin_parts)} 段)"
            )

        seen.add(key)
        rows.append((key, zhuyin, pinyin))

    if seen != title_keys:
        raise ValueError(
            f"{SEARCH_SHEET} 缺少 {len(title_keys - seen)} 个名称 key"
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(SEARCH_EXPORT_HEADER)
        writer.writerows(rows)

    return len(rows)


def main() -> int:
    managed_xlsx = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path("translations_managed.xlsx")
    )
    out_path = (
        Path(sys.argv[2])
        if len(sys.argv) > 2
        else Path("../Localization/search_zh_TW.txt")
    )

    if not managed_xlsx.exists():
        print("错误：找不到翻译管理表：", managed_xlsx)
        return 1

    try:
        wb = load_workbook(managed_xlsx)
        missing = sync_search_sheet(wb)
        wb.save(managed_xlsx)

        if missing:
            print("停止导出：TraditionalChineseSearch 有尚未填写的注音/拼音：")
            for item in missing:
                print(" -", item)
            print("请补全后重新运行导出。")
            return 2

        count = validate_and_export(wb, out_path)
        print(f"繁体中文搜索读音：{count} 条 -> {out_path}")
        return 0
    except PermissionError as exc:
        print("错误：", exc)
        print("请先关闭 Excel/WPS 中打开的 translations_managed.xlsx。")
        return 1
    except Exception as exc:
        print("错误：", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
