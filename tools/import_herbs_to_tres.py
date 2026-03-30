# -*- coding: utf-8 -*-

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook


# ========= 路径 =========
PROJECT_ROOT = Path(r"F:\Doctor")
HERB_XLSX_PATH = Path(r"F:\Doctor\Herb\data\药材.xlsx")
OVERWRITE_EXISTING = True

HERB_DIR = PROJECT_ROOT / "Herb" / "data"
HERB_SCRIPT_RES = "res://Herb/HerbData.gd"


# ========= 数据结构 =========
@dataclass
class HerbRecord:
    herb_id: str
    herb_name: str
    nature: str = ""
    taste: str = ""
    meridians: str = ""
    toxic: str = ""


# ========= 工具 =========
def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)


def cell_str(v) -> str:
    return "" if v is None else str(v).strip()


def safe_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    return name.strip()


def make_id(name: str, explicit: str) -> str:
    return explicit if explicit else re.sub(r"\s+", "_", name)


def quote(text: str) -> str:
    return '"' + text.replace('"', '\\"') + '"'


# ⭐ 核心修复：多分隔符支持
def normalize_multi(text: str, allowed: List[str]) -> str:
    raw = cell_str(text)

    # 统一所有分隔符
    raw = raw.replace("，", ",")
    raw = raw.replace("、", ",")
    raw = raw.replace("/", ",")
    raw = raw.replace(" ", ",")

    result = []
    for x in raw.split(","):
        x = x.strip()
        if x in allowed and x not in result:
            result.append(x)

    return ",".join(result)


# ========= 限定值 =========
ALLOWED_NATURE = ["寒", "热", "温", "凉", "平"]
ALLOWED_TASTE = ["辛", "甘", "酸", "苦", "咸"]
ALLOWED_MERIDIAN = ["表", "心", "肝", "脾", "肺", "肾"]


# ========= 读取 Excel =========
def load_excel(path: Path) -> Dict[str, HerbRecord]:
    wb = load_workbook(path, data_only=True)
    herb_map = {}

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = cell_str(row[0])
            if not name:
                continue

            herb_id = make_id(name, cell_str(row[1]))

            nature = normalize_multi(row[2], ALLOWED_NATURE)
            taste = normalize_multi(row[3], ALLOWED_TASTE)
            meridians = normalize_multi(row[4], ALLOWED_MERIDIAN)
            toxic = cell_str(row[5])

            herb_map[name] = HerbRecord(
                herb_id=herb_id,
                herb_name=name,
                nature=nature,
                taste=taste,
                meridians=meridians,
                toxic=toxic,
            )

    return herb_map


# ========= 写 tres =========
def to_tres(h: HerbRecord) -> str:
    return f"""[gd_resource type="Resource" script_class="HerbData" load_steps=2 format=3]

[ext_resource type="Script" path="{HERB_SCRIPT_RES}" id="1"]

[resource]
script = ExtResource("1")
herb_id = {quote(h.herb_id)}
herb_name = {quote(h.herb_name)}
nature = {quote(h.nature)}
taste = {quote(h.taste)}
meridians = {quote(h.meridians)}
toxic = {quote(h.toxic)}
"""


def write_all(herb_map: Dict[str, HerbRecord]):
    ensure_dir(HERB_DIR)

    for h in herb_map.values():
        path = HERB_DIR / f"{safe_filename(h.herb_name)}.tres"

        if path.exists() and not OVERWRITE_EXISTING:
            continue

        path.write_text(to_tres(h), encoding="utf-8")


# ========= 主入口 =========
def main():
    print("==== 开始导入药材 ====")

    herb_map = load_excel(HERB_XLSX_PATH)
    print(f"读取 {len(herb_map)} 条")

    write_all(herb_map)

    print("==== 完成 → res://Herb/data/ ====")


if __name__ == "__main__":
    main()