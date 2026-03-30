# -*- coding: utf-8 -*-

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


# ========= 路径 =========
PROJECT_ROOT = Path(r"F:\Doctor")
FORMULA_XLSX_PATH = Path(r"F:\Doctor\Formula\data\方剂.xlsx")
OVERWRITE_EXISTING = True

FORMULA_DIR = PROJECT_ROOT / "Formula" / "data"
FORMULA_SCRIPT_RES = "res://Formula/FormulaData.gd"
FORMULA_INGREDIENT_SCRIPT_RES = "res://Formula/FormulaIngredient.gd"


# ========= 数据结构 =========
@dataclass
class IngredientRecord:
    herb_name: str
    role: str
    amount: float
    unit: str


@dataclass
class FormulaRecord:
    formula_id: str
    formula_name: str
    source_book: str = ""
    target_disease_id: str = ""
    target_disease_name: str = ""
    jun_group: List[IngredientRecord] = field(default_factory=list)
    chen_group: List[IngredientRecord] = field(default_factory=list)
    zuo_group: List[IngredientRecord] = field(default_factory=list)
    shi_group: List[IngredientRecord] = field(default_factory=list)


# ========= 工具 =========
def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)


def cell_str(v) -> str:
    return "" if v is None else str(v).strip()


def safe_filename(name: str) -> str:
    name = cell_str(name)
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    return name.strip() if name.strip() else "unnamed"


def make_id(name: str) -> str:
    text = cell_str(name)
    text = re.sub(r"\s+", "_", text)
    text = text.replace("-", "_")
    text = re.sub(r"[^\w\u4e00-\u9fff_]+", "", text)
    return text if text else "unnamed_id"


def quote(text: str) -> str:
    text = "" if text is None else str(text)
    text = text.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")
    return f'"{text}"'


def normalize_role(role: str) -> str:
    role = cell_str(role)
    if role in ["君", "臣", "佐", "使"]:
        return role
    return "佐"


# ========= 单位换算 =========
UNIT_TO_FEN = {
    "分": 1,
    "钱": 10,
    "两": 100,
    "斤": 1600,
}


def parse_amount_text(amount_text: str) -> Tuple[float, str]:
    """
    支持：
    - 3钱
    - 1两
    - 1钱5分
    - 1两2钱
    最终优先转成：
    - 能整除斤 -> jin
    - 能整除两 -> liang
    - 能整除钱 -> qian
    - 否则 -> fen
    """
    text = cell_str(amount_text)
    if not text:
        return 0.0, "qian"

    text = text.replace(" ", "").replace("　", "")
    parts = re.findall(r"(\d+(?:\.\d+)?)(分|钱|两|斤)", text)

    if not parts:
        return 0.0, "qian"

    total_fen = 0.0
    for number_text, unit_text in parts:
        total_fen += float(number_text) * UNIT_TO_FEN[unit_text]

    if total_fen % 1600 == 0:
        return total_fen / 1600, "jin"
    if total_fen % 100 == 0:
        return total_fen / 100, "liang"
    if total_fen % 10 == 0:
        return total_fen / 10, "qian"

    return total_fen, "fen"


# ========= 读取 Excel =========
def load_formulas(path: Path) -> Dict[str, FormulaRecord]:
    """
    Excel 列顺序：
    0 方剂ID
    1 方剂名
    2 药材
    3 角色
    4 用量
    5 疾病ID
    6 主治
    7 来源
    """
    wb = load_workbook(path, data_only=True)
    formula_map: Dict[str, FormulaRecord] = {}

    for ws in wb.worksheets:
        current_formula: Optional[FormulaRecord] = None

        for row in ws.iter_rows(min_row=2, values_only=True):
            formula_id = cell_str(row[0] if len(row) > 0 else "")
            formula_name = cell_str(row[1] if len(row) > 1 else "")
            herb_name = cell_str(row[2] if len(row) > 2 else "")
            role = cell_str(row[3] if len(row) > 3 else "")
            amount_text = cell_str(row[4] if len(row) > 4 else "")
            disease_id = cell_str(row[5] if len(row) > 5 else "")
            disease_name = cell_str(row[6] if len(row) > 6 else "")
            source_book = cell_str(row[7] if len(row) > 7 else "")

            if not any(cell_str(v) for v in row):
                continue

            # 新方剂开始：以“方剂名”是否有值为准
            if formula_name:
                current_formula = FormulaRecord(
                    formula_id=formula_id if formula_id else make_id(formula_name),
                    formula_name=formula_name,
                    source_book=source_book if source_book else ws.title,
                    target_disease_id=disease_id,
                    target_disease_name=disease_name,
                )
                formula_map[formula_name] = current_formula

            if current_formula is None:
                continue

            if herb_name:
                amount, unit = parse_amount_text(amount_text)
                ingredient = IngredientRecord(
                    herb_name=herb_name,
                    role=normalize_role(role),
                    amount=amount,
                    unit=unit,
                )

                if ingredient.role == "君":
                    current_formula.jun_group.append(ingredient)
                elif ingredient.role == "臣":
                    current_formula.chen_group.append(ingredient)
                elif ingredient.role == "佐":
                    current_formula.zuo_group.append(ingredient)
                else:
                    current_formula.shi_group.append(ingredient)

    return formula_map


# ========= 写 tres =========
def role_group_field_name(role: str) -> str:
    if role == "君":
        return "jun_group"
    if role == "臣":
        return "chen_group"
    if role == "佐":
        return "zuo_group"
    return "shi_group"


def formula_to_tres_text(formula: FormulaRecord) -> str:
    ext_lines: List[str] = []
    sub_lines: List[str] = []

    # 固定脚本引用
    ext_lines.append(
        f'[ext_resource type="Script" path="{FORMULA_INGREDIENT_SCRIPT_RES}" id="1_ingredient_script"]'
    )
    ext_lines.append(
        f'[ext_resource type="Script" path="{FORMULA_SCRIPT_RES}" id="2_formula_script"]'
    )

    herb_ext_id_map: Dict[str, str] = {}
    group_to_sub_ids = {
        "jun_group": [],
        "chen_group": [],
        "zuo_group": [],
        "shi_group": [],
    }

    def get_herb_ext_id(herb_name: str) -> str:
        if herb_name in herb_ext_id_map:
            return herb_ext_id_map[herb_name]

        herb_path = f"res://Herb/data/{safe_filename(herb_name)}.tres"
        ext_id = f"herb_{len(herb_ext_id_map) + 1}"
        ext_lines.append(
            f'[ext_resource type="Resource" path="{herb_path}" id="{ext_id}"]'
        )
        herb_ext_id_map[herb_name] = ext_id
        return ext_id

    ingredient_index = 1

    def add_ingredient(ing: IngredientRecord):
        nonlocal ingredient_index

        herb_ext_id = get_herb_ext_id(ing.herb_name)
        sub_id = f"Ingredient_{ingredient_index}"
        ingredient_index += 1

        sub_lines.extend([
            f'[sub_resource type="Resource" id="{sub_id}"]',
            'script = ExtResource("1_ingredient_script")',
            f'herb = ExtResource("{herb_ext_id}")',
            f"amount = {float(ing.amount)}",
            f"unit = {quote(ing.unit)}",
            "required = true",
            "",
        ])

        field = role_group_field_name(ing.role)
        group_to_sub_ids[field].append(sub_id)

    for ing in formula.jun_group:
        add_ingredient(ing)
    for ing in formula.chen_group:
        add_ingredient(ing)
    for ing in formula.zuo_group:
        add_ingredient(ing)
    for ing in formula.shi_group:
        add_ingredient(ing)

    def group_text(sub_ids: List[str]) -> str:
        inner = ", ".join(f'SubResource("{sid}")' for sid in sub_ids)
        return f"[{inner}]"

    load_steps = 2 + len(herb_ext_id_map) + sum(len(v) for v in group_to_sub_ids.values())

    lines: List[str] = [
        f'[gd_resource type="Resource" script_class="FormulaData" load_steps={load_steps} format=3]',
        "",
    ]
    lines.extend(ext_lines)
    lines.append("")
    lines.extend(sub_lines)
    lines.extend([
        "[resource]",
        'script = ExtResource("2_formula_script")',
        f"formula_id = {quote(formula.formula_id)}",
        f"formula_name = {quote(formula.formula_name)}",
        f"source_book = {quote(formula.source_book)}",
        f"target_disease_id = {quote(formula.target_disease_id)}",
        f"target_disease_name = {quote(formula.target_disease_name)}",
        f'effect_text = ""',
        f"indication_text = {quote(formula.target_disease_name)}",
        f'description = ""',
        f'jun_group = {group_text(group_to_sub_ids["jun_group"])}',
        f'chen_group = {group_text(group_to_sub_ids["chen_group"])}',
        f'zuo_group = {group_text(group_to_sub_ids["zuo_group"])}',
        f'shi_group = {group_text(group_to_sub_ids["shi_group"])}',
        "",
    ])

    return "\n".join(lines)


def write_all(formula_map: Dict[str, FormulaRecord]):
    ensure_dir(FORMULA_DIR)

    for formula in formula_map.values():
        source_folder = FORMULA_DIR / safe_filename(formula.source_book)
        ensure_dir(source_folder)

        path = source_folder / f"{safe_filename(formula.formula_name)}.tres"

        if path.exists() and not OVERWRITE_EXISTING:
            continue

        path.write_text(formula_to_tres_text(formula), encoding="utf-8")


# ========= 主入口 =========
def main():
    print("==== 开始导入方剂 ====")

    formula_map = load_formulas(FORMULA_XLSX_PATH)
    print(f"读取 {len(formula_map)} 条方剂")

    write_all(formula_map)

    print("==== 完成 → res://Formula/data/ ====")


if __name__ == "__main__":
    main()